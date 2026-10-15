"""Reads back the xlsx written by scripts/check-leads-export.mjs and asserts text, types and styling.
Usage: pip install openpyxl; python scripts/check-leads-export.py sample.xlsx
"""
import sys
import zipfile
import xml.dom.minidom
from datetime import datetime

from openpyxl import load_workbook

path = sys.argv[1]
failed = 0
total = 0


def check(name, want, got):
    global failed, total
    total += 1
    if want != got:
        failed += 1
        print(f"FAIL {name}\n  want: {want!r}\n  got:  {got!r}")


# Every XML part is well formed; parts that carry a prolog declare UTF-8 (sharedStrings keeps the bare default).
with zipfile.ZipFile(path) as z:
    for n in z.namelist():
        if n.endswith(('.xml', '.rels')):
            raw = z.read(n)
            xml.dom.minidom.parseString(raw)
            head = raw[:80].decode('utf-8')
            if n != 'xl/sharedStrings.xml':
                check(f"utf-8 prolog {n}", True, 'encoding="UTF-8"' in head)
    check("no formula elements in sheet", False, '<f>' in z.read('xl/worksheets/sheet1.xml').decode('utf-8'))
    check("no inline strings", False, 't="inlineStr"' in z.read('xl/worksheets/sheet1.xml').decode('utf-8'))

wb = load_workbook(path)
check("sheet name", ["Khách hàng"], wb.sheetnames)
ws = wb.active
check("freeze panes", "A2", ws.freeze_panes)
check("auto filter", "A1:O7", ws.auto_filter.ref)

headers = [c.value for c in ws[1]]
check("headers", ["STT", "Thời gian", "Nguồn", "Họ và tên", "Số điện thoại", "Email", "Chủ đề", "Nhu cầu", "Khu vực",
                  "Loại BĐS", "Mức giá", "Ghi chú của khách", "Trạng thái", "Ghi chú nội bộ", "Tin ký gửi liên kết"], headers)
h = ws["A1"]
check("header bold", True, bool(h.font.b))
check("header font colour", "FFFFFFFF", h.font.color.rgb)
check("header fill", "FF0F2942", h.fill.fgColor.rgb)
check("header font", ("Calibri", 11), (h.font.name, h.font.sz))
check("body font", ("Calibri", 11), (ws["D2"].font.name, ws["D2"].font.sz))

widths = [round(ws.column_dimensions[c].width) for c in "ABCDEFGHIJKLMNO"]
check("column widths", [6, 17, 10, 24, 16, 28, 28, 11, 22, 18, 18, 60, 16, 40, 34], widths)
check("max width <= 60", True, max(widths) <= 60)
check("wrap on note", True, bool(ws["L2"].alignment.wrap_text))

col = {name: i + 1 for i, name in enumerate(headers)}
val = lambda r, name: ws.cell(row=r, column=col[name])

check("name round-trip", "Nguyễn Thị Hương", val(2, "Họ và tên").value)
check("name 2 round-trip", "Đặng Quốc Việt", val(3, "Họ và tên").value)
check("emoji + newline", "Cần bán gấp 🏠🙂\nDòng hai", val(2, "Ghi chú của khách").value)
check("phone string", ("0931123416", "s", "@"), (val(2, "Số điện thoại").value, val(2, "Số điện thoại").data_type, val(2, "Số điện thoại").number_format))
check("phone +84", ("+84 931 123 416", "s"), (val(3, "Số điện thoại").value, val(3, "Số điện thoại").data_type))
check("phone 0123456789", ("0123456789", "s"), (val(6, "Số điện thoại").value, val(6, "Số điện thoại").data_type))

# Typed date, GMT+7 wall clock
d = val(2, "Thời gian")
check("date typed", (datetime(2026, 10, 15, 14, 30), "d", "dd/mm/yyyy hh:mm"), (d.value, d.data_type, d.number_format))
check("invalid date blank", None, val(5, "Thời gian").value)

# Formula injection: plain strings with the text format, nothing evaluated
for r, name, text in [
    (4, "Họ và tên", '=HYPERLINK("http://evil.example","Bấm vào")'),
    (4, "Số điện thoại", "=1+1"),
    (4, "Ghi chú của khách", "+cmd|' /C calc'!A0"),
    (4, "Ghi chú nội bộ", "@SUM(1,1)"),
    (5, "Họ và tên", "-2+3"),
]:
    c = val(r, name)
    check(f"injection r{r} {name}", (text, "s", "@"), (c.value, c.data_type, c.number_format))
    check(f"injection r{r} {name} not formula", False, c.data_type == "f")

check("contact email", "viet@example.com", val(3, "Email").value)
check("contact subject", "Tư vấn dự án BĐS", val(3, "Chủ đề").value)
check("long note capped", 32001, len(val(5, "Ghi chú của khách").value))
check("illegal chars stripped", "ABC", val(7, "Họ và tên").value)
check("row height for tall row", True, (ws.row_dimensions[5].height or 0) >= 150)
check("rows", 7, ws.max_row)

print(f"{total - failed}/{total} checks passed")
sys.exit(1 if failed else 0)
